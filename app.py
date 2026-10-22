from flask import Flask, request, jsonify, send_from_directory, send_file, session
from werkzeug.security import generate_password_hash, check_password_hash
import requests, json, os, re, sqlite3, threading, webbrowser, unicodedata, time, secrets, shutil
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from functools import wraps

IS_RENDER = os.getenv("RENDER", "").lower() == "true" or bool(os.getenv("RENDER_SERVICE_ID"))
app = Flask(__name__, static_folder="static")
app.secret_key = os.getenv("SECRET_KEY") or "keea-opinion-map-local-secret"
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=IS_RENDER,
)

# AI_PROVIDER=ollama: 로컬/원격 Ollama 서버 사용
# AI_PROVIDER=openai: OpenAI 또는 OpenAI 호환 API 사용
AI_PROVIDER = os.getenv("AI_PROVIDER", "ollama").strip().lower()
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434/api/chat")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:3b")
AI_BASE_URL = os.getenv("AI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
AI_API_KEY = os.getenv("AI_API_KEY") or os.getenv("OPENAI_API_KEY", "")
AI_MODEL = os.getenv("AI_MODEL", "gpt-4o-mini")
AI_TIMEOUT = int(os.getenv("AI_TIMEOUT", "90"))
ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin1234")

if IS_RENDER and (not os.getenv("ADMIN_USERNAME") or not os.getenv("ADMIN_PASSWORD")):
    raise RuntimeError("Render 환경변수 ADMIN_USERNAME과 ADMIN_PASSWORD를 반드시 설정하세요.")

def persistent_db_path():
    configured=os.getenv("DB_PATH")
    if configured: return os.path.abspath(os.path.expanduser(configured))
    if os.name=="nt": base=os.getenv("LOCALAPPDATA") or os.getenv("APPDATA") or os.path.expanduser("~")
    elif sys_platform()=="darwin": base=os.path.join(os.path.expanduser("~"),"Library","Application Support")
    else: base=os.getenv("XDG_DATA_HOME") or os.path.join(os.path.expanduser("~"),".local","share")
    return os.path.join(base,"KEEA-AI-Opinion-Map","opinion_map.db")

def sys_platform():
    import sys
    return sys.platform

DB_PATH = persistent_db_path()

def prepare_db_storage():
    parent=os.path.dirname(DB_PATH)
    if parent: os.makedirs(parent,exist_ok=True)
    legacy=os.path.join(os.path.dirname(__file__),"opinion_map.db")
    if not os.path.exists(DB_PATH) and os.path.abspath(legacy)!=os.path.abspath(DB_PATH) and os.path.exists(legacy):
        shutil.copy2(legacy,DB_PATH)

@app.after_request
def disable_browser_cache(response):
    response.headers["Cache-Control"]="no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"]="no-cache"
    response.headers["Expires"]="0"
    return response

def db():
    c=sqlite3.connect(DB_PATH); c.row_factory=sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); return c

def init_db():
    prepare_db_storage()
    with db() as c: c.executescript("""
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,username TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,is_admin INTEGER NOT NULL DEFAULT 0);
    CREATE TABLE IF NOT EXISTS nodes(id TEXT PRIMARY KEY,kind TEXT NOT NULL,text TEXT NOT NULL,normalized TEXT NOT NULL,title TEXT NOT NULL,category TEXT NOT NULL,keywords TEXT NOT NULL DEFAULT '[]',related TEXT NOT NULL DEFAULT '[]',parent_id TEXT,display_size TEXT NOT NULL DEFAULT 'medium',created_by INTEGER NOT NULL REFERENCES users(id),created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS contributions(node_id TEXT NOT NULL REFERENCES nodes(id) ON DELETE CASCADE,user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,text TEXT NOT NULL,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,PRIMARY KEY(node_id,user_id));
    """)
    with db() as c:
        if "is_admin" not in [r[1] for r in c.execute("PRAGMA table_info(users)")]: c.execute("ALTER TABLE users ADD COLUMN is_admin INTEGER NOT NULL DEFAULT 0")
        if "parent_id" not in [r[1] for r in c.execute("PRAGMA table_info(nodes)")]: c.execute("ALTER TABLE nodes ADD COLUMN parent_id TEXT")
        if "display_size" not in [r[1] for r in c.execute("PRAGMA table_info(nodes)")]: c.execute("ALTER TABLE nodes ADD COLUMN display_size TEXT NOT NULL DEFAULT 'medium'")
        c.execute("UPDATE nodes SET parent_id=NULL WHERE kind='question'")
        kinds={r["id"]:r["kind"] for r in c.execute("SELECT id,kind FROM nodes").fetchall()}
        for r in c.execute("SELECT id,kind,related FROM nodes").fetchall():
            rel=json.loads(r["related"])
            cleaned=[x for x in rel if not (r["kind"]=="question" and kinds.get(x)=="question")]
            if cleaned!=rel: c.execute("UPDATE nodes SET related=? WHERE id=?",(json.dumps(cleaned),r["id"]))
        admin=c.execute("SELECT id FROM users WHERE username=?",(ADMIN_USERNAME,)).fetchone()
        if admin: c.execute("UPDATE users SET password_hash=?,is_admin=1 WHERE id=?",(generate_password_hash(ADMIN_PASSWORD),admin["id"]))
        else: c.execute("INSERT INTO users(username,password_hash,is_admin) VALUES(?,?,1)",(ADMIN_USERNAME,generate_password_hash(ADMIN_PASSWORD)))

def norm(s): return re.sub(r"\s+"," ",unicodedata.normalize("NFKC",str(s or "")).strip()).lower()

def call_ai_json(system_prompt, user_prompt, temperature=0.0):
    """JSON 응답을 반환하는 AI 호출 함수. Ollama와 OpenAI 호환 API를 지원한다."""
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]

    if AI_PROVIDER in ("openai", "openai-compatible", "openai_compatible"):
        if not AI_API_KEY:
            raise RuntimeError("AI_API_KEY 환경변수가 없습니다. Render 환경변수에 AI API 키를 설정하세요.")
        endpoint = AI_BASE_URL
        if not endpoint.endswith("/chat/completions"):
            endpoint += "/chat/completions"
        response = requests.post(
            endpoint,
            headers={"Authorization": f"Bearer {AI_API_KEY}"},
            json={
                "model": AI_MODEL,
                "messages": messages,
                "temperature": temperature,
                "response_format": {"type": "json_object"},
            },
            timeout=AI_TIMEOUT,
        )
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"]
    elif AI_PROVIDER == "ollama":
        response = requests.post(
            OLLAMA_URL,
            json={
                "model": OLLAMA_MODEL,
                "messages": messages,
                "stream": False,
                "format": "json",
                "options": {"temperature": temperature},
            },
            timeout=AI_TIMEOUT,
        )
        response.raise_for_status()
        content = response.json()["message"]["content"]
    else:
        raise ValueError("AI_PROVIDER는 ollama 또는 openai로 설정해야 합니다.")

    result = json.loads(content)
    if not isinstance(result, dict):
        raise ValueError("AI 응답이 JSON 객체가 아닙니다.")
    return result

def policy_families(text):
    t=norm(text)
    return {name for name,terms in POLICY_FAMILIES.items() if any(norm(term) in t for term in terms)}

def hard_policy_conflict(a,b):
    fa,fb=policy_families(a),policy_families(b)
    return bool(fa and fb and fa.isdisjoint(fb))

def policy_candidate_text(candidate):
    parts=[candidate.get("title") or "",candidate.get("text") or "",candidate.get("category") or ""]
    parts.extend(candidate.get("keywords") or [])
    for item in candidate.get("opinions") or []:
        parts.append(item.get("text") if isinstance(item,dict) else str(item))
    return " ".join(str(x) for x in parts if x)

def choose_policy_merge_id(source_texts,candidates,parent=None):
    source_texts=[str(x).strip() for x in source_texts if str(x or "").strip()]
    if not source_texts or not candidates: return None
    source_joined=" ".join(source_texts)
    source_norms={norm(x) for x in source_texts if norm(x)}

    # 완전히 같은 제목/원문이 있으면 AI를 거치지 않고 기존 노드로 합친다.
    for c in candidates:
        candidate_values=[c.get("title") or "",c.get("text") or ""]+[x.get("text","") for x in (c.get("opinions") or []) if isinstance(x,dict)]
        if source_norms.intersection({norm(x) for x in candidate_values if norm(x)}):
            return str(c["id"])

    # 교통비·유연근무·육아휴직 등 명시적으로 정의한 정책군은 로컬 AI의 오판보다
    # 규칙을 우선한다. 같은 단일 정책군 후보가 여러 개면 참여자가 많은 기존 노드로 수렴시킨다.
    source_families=policy_families(source_joined)
    if len(source_families)==1:
        family=next(iter(source_families))
        family_matches=[]
        for c in candidates:
            cf=policy_families(policy_candidate_text(c))
            if cf=={family}: family_matches.append(c)
        if family_matches:
            best=max(family_matches,key=lambda c:len(c.get("opinions") or []))
            return str(best["id"])

    safe_candidates=[c for c in candidates if not hard_policy_conflict(source_joined,policy_candidate_text(c))]
    if not safe_candidates: return None
    compact=[{"id":c["id"],"title":c.get("title"),"category":c.get("category"),"keywords":c.get("keywords"),"opinions":c.get("opinions") or [{"text":c.get("text","")}]} for c in safe_candidates]
    try:
        prompt=f"""선택된 공지: {json.dumps(parent,ensure_ascii=False)}
새 의견 또는 재분류된 의견 묶음: {json.dumps(source_texts,ensure_ascii=False)}
같은 공지 아래의 기존 의견 후보: {json.dumps(compact,ensure_ascii=False)}

'관련 있음'이 아니라 정확히 같은 제도·정책 대상을 다루는 후보 하나만 고른다.
- 표현, 범위, 대상, 시행 방식만 다르고 정책 대상이 같으면 같은 묶음이다.
- 교통비/통근비 지원과 유연근무제는 서로 다른 정책이므로 절대 합치지 않는다.
- 임금과 성과급, 유연근무와 육아휴직처럼 제도 자체가 다르면 절대 합치지 않는다.
- 확신할 수 없으면 반드시 null을 반환한다.
JSON만 반환: {{"merge_id":"후보 id 또는 null"}}"""
        result = call_ai_json(
            "사내 의견을 정책 대상 단위로 매우 엄격하게 병합한다. 관련성만으로는 병합하지 않는다. JSON만 출력한다.",
            prompt,
            temperature=0,
        )
        chosen=str(result.get("merge_id")) if result.get("merge_id") else None
        candidate=next((x for x in safe_candidates if str(x["id"])==chosen),None)
        if not candidate: return None
        candidate_joined=policy_candidate_text(candidate)
        if hard_policy_conflict(source_joined,candidate_joined): return None
        verify=f"""새 의견 묶음: {json.dumps(source_texts,ensure_ascii=False)}
기존 의견 묶음: {json.dumps(candidate,ensure_ascii=False)}

두 묶음이 정확히 동일한 제도·정책 대상을 다루는지 판정한다.
교통비/통근비 지원과 유연근무제는 false다.
임금과 성과급도 false다. 단지 같은 공지나 복지·근무 분야라는 이유로 true로 하지 않는다.
표현이나 세부 범위만 다르고 정책 대상 자체가 같을 때만 true다.
JSON만 반환: {{"same_policy":true 또는 false}}"""
        verdict = call_ai_json(
            "정책 대상 동일성을 보수적으로 판정한다. 애매하면 false. JSON만 출력한다.",
            verify,
            temperature=0,
        )
        return chosen if verdict.get("same_policy") is True else None
    except Exception:
        return None

def load_policy_candidates(parent_id,exclude_ids=()):
    exclude_ids={str(x) for x in exclude_ids}
    with db() as c:
        if parent_id is None:
            rows=c.execute("SELECT * FROM nodes WHERE kind='opinion' AND parent_id IS NULL ORDER BY created_at,id").fetchall()
        else:
            rows=c.execute("SELECT * FROM nodes WHERE kind='opinion' AND parent_id=? ORDER BY created_at,id",(parent_id,)).fetchall()
        result=[]
        for r in rows:
            if str(r["id"]) in exclude_ids: continue
            opinions=[dict(x) for x in c.execute("SELECT co.user_id,u.username author,co.text FROM contributions co JOIN users u ON u.id=co.user_id WHERE co.node_id=? ORDER BY co.created_at,co.user_id",(r["id"],)).fetchall()]
            result.append({"id":r["id"],"text":r["text"],"title":r["title"],"category":r["category"],"keywords":json.loads(r["keywords"]),"opinions":opinions,"parent_id":r["parent_id"]})
        return result

def move_contribution(c,source_id,target_id,user_id,text):
    existing=c.execute("SELECT text FROM contributions WHERE node_id=? AND user_id=?",(target_id,user_id)).fetchone()
    if existing:
        c.execute("DELETE FROM contributions WHERE node_id=? AND user_id=?",(source_id,user_id))
        return 1
    moved=c.execute("UPDATE contributions SET node_id=? WHERE node_id=? AND user_id=?",(target_id,source_id,user_id))
    if moved.rowcount!=1: raise RuntimeError("재분류 중 원본 의견이 변경되었습니다.")
    return 0

def session_user():
    uid=session.get("user_id")
    if not uid: return None
    with db() as c: user=c.execute("SELECT id,username,is_admin FROM users WHERE id=?",(uid,)).fetchone()
    if not user:
        session.clear()
        return None
    session["username"]=user["username"]
    session["is_admin"]=bool(user["is_admin"])
    return user

def login_required(fn):
    @wraps(fn)
    def wrapped(*a,**k):
        if not session_user(): return jsonify(error="로그인이 만료되었습니다. 다시 로그인하거나 새 계정을 만드세요."),401
        return fn(*a,**k)
    return wrapped

def node_rows(c):
    admin=bool(session.get("is_admin"))
    rows=c.execute("""SELECT n.*,u.username,CASE WHEN n.kind='opinion' THEN COUNT(co.user_id) ELSE 1 END support_count,MAX(CASE WHEN co.user_id=? THEN 1 ELSE 0 END) supported_by_me FROM nodes n JOIN users u ON u.id=n.created_by LEFT JOIN contributions co ON co.node_id=n.id GROUP BY n.id ORDER BY n.created_at,n.id""",(session["user_id"],)).fetchall()
    result=[]
    for r in rows:
        opinions=[]
        if r["kind"]=="opinion": opinions=[dict(x) for x in c.execute("SELECT co.user_id,u.username author,co.text FROM contributions co JOIN users u ON u.id=co.user_id WHERE co.node_id=? ORDER BY co.created_at",(r["id"],)).fetchall()]
        can_delete=admin or (bool(r["supported_by_me"]) if r["kind"]=="opinion" else r["created_by"]==session["user_id"])
        delete_mode="all" if admin or r["kind"]=="question" else "own"
        result.append({"id":r["id"],"kind":r["kind"],"text":r["text"],"title":r["title"],"category":r["category"],"keywords":json.loads(r["keywords"]),"related":json.loads(r["related"]),"parent_id":r["parent_id"],"display_size":r["display_size"],"author":r["username"],"support_count":r["support_count"],"opinions":opinions,"supported_by_me":bool(r["supported_by_me"]),"can_delete":can_delete,"delete_mode":delete_mode,"can_resize":r["kind"]=="question" and (admin or r["created_by"]==session["user_id"])})
    return result

@app.get("/")
def index():
    return send_from_directory(os.path.join(os.path.dirname(__file__), "static"), "index.html")

@app.get("/health")
def health():
    return jsonify(status="ok")

@app.post("/api/register")
def register():
    d=request.get_json() or {}; username=(d.get("username") or "").strip(); password=d.get("password") or ""
    if len(username)<2 or len(password)<4: return jsonify(error="이름은 2자, 비밀번호는 4자 이상 입력하세요."),400
    try:
        with db() as c:
            cur=c.execute("INSERT INTO users(username,password_hash) VALUES(?,?)",(username,generate_password_hash(password)))
            session["user_id"]=cur.lastrowid; session["username"]=username
        return jsonify(username=username,is_admin=False)
    except sqlite3.IntegrityError: return jsonify(error="이미 사용 중인 이름입니다."),409

@app.post("/api/login")
def login():
    d=request.get_json() or {}; username=(d.get("username") or "").strip(); password=d.get("password") or ""
    with db() as c: u=c.execute("SELECT * FROM users WHERE username=?",(username,)).fetchone()
    if not u or not check_password_hash(u["password_hash"],password): return jsonify(error="이름 또는 비밀번호가 맞지 않습니다."),401
    session["user_id"]=u["id"]; session["username"]=u["username"]; session["is_admin"]=bool(u["is_admin"]); return jsonify(username=u["username"],is_admin=bool(u["is_admin"]))

@app.post("/api/logout")
def logout(): session.clear(); return jsonify(ok=True)

@app.get("/api/me")
def me():
    user=session_user()
    return jsonify(username=user["username"],is_admin=bool(user["is_admin"])) if user else (jsonify(error="로그인이 필요합니다."),401)

@app.get("/api/nodes")
@login_required
def nodes():
    with db() as c: return jsonify(node_rows(c))

@app.post("/api/analyze")
@login_required
def analyze():
    d=request.get_json() or {}; text=(d.get("text") or "").strip(); kind=d.get("kind","opinion"); parent=d.get("parent_question") or None
    if not text: return jsonify(error="내용을 입력해 주세요."),400
    with db() as c: existing=node_rows(c)
    prompt=f"""종류: {kind}\n내용: {text}\n선택된 공지: {json.dumps(parent,ensure_ascii=False)}\n기존 노트: {json.dumps(existing,ensure_ascii=False)}
JSON만 반환:
{{"title":"하나의 핵심 키워드 또는 짧은 문장","category":"의미 분류명","keywords":["핵심어"],"related_ids":["관련 노트 id"]}}
규칙:
1) question은 공지 노트이며 제목을 5~16자로 요약.
2) opinion 제목은 입력 의견 자체의 문장 또는 핵심어로 작성하고 공지 제목을 복사하지 않는다.
3) 사용자가 선택한 공지와의 연결 여부는 시스템이 처리하므로 판단하거나 거부하지 않는다.
4) 공지 노트끼리는 절대 연결하지 않는다. 다른 question을 related_ids에 넣지 않는다.
5) 노트 표시는 문장 전체보다 핵심 키워드와 짧은 요약을 우선한다. 같은 의미 범주는 동일한 category명을 사용.
"""
    try:
        result = call_ai_json(
            "한국어 사내 공지와 의견을 의미 기반으로 분류한다. JSON만 출력한다.",
            prompt,
            temperature=0.1,
        )
        if not isinstance(result,dict): raise ValueError()
        merge_id=None
        parent_id=str(parent.get("id")) if isinstance(parent,dict) and parent.get("id") else None
        candidates=[{"id":x["id"],"text":x["text"],"title":x["title"],"category":x["category"],"keywords":x["keywords"],"opinions":x.get("opinions") or [{"text":x["text"]}]} for x in existing if x["kind"]=="opinion" and x.get("parent_id")==parent_id]
        if kind=="opinion" and candidates:
            source=[text,str(result.get("title") or "")]+[str(x) for x in (result.get("keywords") if isinstance(result.get("keywords"),list) else [])]
            merge_id=choose_policy_merge_id(source,candidates,parent)
        return jsonify(title=str(result.get("title") or text)[:40],category=str(result.get("category") or "기타")[:30],keywords=[str(x)[:30] for x in (result.get("keywords") if isinstance(result.get("keywords"),list) else [])[:8]],related_ids=[str(x) for x in (result.get("related_ids") if isinstance(result.get("related_ids"),list) else [])[:4]],merge_id=merge_id)
    except Exception:
        words=[x for x in re.sub(r"[^\w가-힣 ]"," ",text).split() if len(x)>1]
        return jsonify(title=(" ".join(words[:3]) or text)[:16],category=(words[0] if words else "기타"),keywords=words[:4],related_ids=[],merge_id=None)

@app.post("/api/nodes/<node_id>/reclassify")
@login_required
def reclassify_bundle(node_id):
    with db() as c:
        node=c.execute("SELECT * FROM nodes WHERE id=? AND kind='opinion'",(node_id,)).fetchone()
        if not node: return jsonify(error="의견 노트를 찾을 수 없습니다."),404
        rows=[dict(x) for x in c.execute("SELECT co.user_id,u.username author,co.text FROM contributions co JOIN users u ON u.id=co.user_id WHERE co.node_id=? ORDER BY co.created_at,co.user_id",(node_id,)).fetchall()]
    if not rows: return jsonify(error="재분류할 의견이 없습니다."),400

    # 이미 분리되어 1명짜리 노드가 된 경우에도 같은 공지의 기존 정책 노드를 찾아 다시 병합한다.
    if len(rows)==1:
        candidates=load_policy_candidates(node["parent_id"],{node_id})
        source=[rows[0]["text"],node["title"]]+json.loads(node["keywords"])
        target=choose_policy_merge_id(source,candidates,{"id":node["parent_id"]} if node["parent_id"] else None)
        if not target:
            return jsonify(ok=True,groups=1,message="다른 동일 정책 노드를 찾지 못했습니다.")
        try:
            with db() as c:
                c.execute("BEGIN IMMEDIATE")
                current=c.execute("SELECT user_id,text FROM contributions WHERE node_id=?",(node_id,)).fetchall()
                if len(current)!=1 or current[0]["user_id"]!=rows[0]["user_id"] or current[0]["text"]!=rows[0]["text"]:
                    return jsonify(error="의견이 이미 변경되었습니다. 화면을 새로고침해 주세요."),409
                dropped=move_contribution(c,node_id,target,rows[0]["user_id"],rows[0]["text"])
                c.execute("DELETE FROM nodes WHERE id=?",(node_id,))
                for r in c.execute("SELECT id,related FROM nodes").fetchall():
                    rel=[x for x in json.loads(r["related"]) if x!=node_id]
                    c.execute("UPDATE nodes SET related=? WHERE id=?",(json.dumps(rel),r["id"]))
            return jsonify(ok=True,groups=1,merged_into=target,dropped_duplicate=dropped,message="기존 동일 정책 의견 노드로 병합했습니다.")
        except (sqlite3.Error,RuntimeError):
            return jsonify(error="기존 의견 노드로 병합하는 중 데이터가 변경되었습니다. 다시 시도해 주세요."),409

    items=[{"ref":f"c{i}","author":x["author"],"text":x["text"]} for i,x in enumerate(rows)]
    prompt=f"""개별 사내 의견: {json.dumps(items,ensure_ascii=False)}
동일한 제도·정책 대상끼리만 그룹으로 묶어라.
- 출퇴근 유연제 확대와 유연제 대상 확대는 같은 그룹.
- 교통비/통근비 지원과 유연근무제는 서로 다른 그룹.
- 출퇴근 유연제, 출퇴근 보조비, 육아휴직은 각각 다른 그룹.
- 같은 공지에 속하거나 관련 분야라는 이유만으로 묶지 않는다.
- 모든 ref를 정확히 한 번씩 포함한다.
JSON만 반환: {{"groups":[{{"title":"짧은 제목","category":"분류","keywords":["핵심어"],"refs":["c0"]}}]}}"""
    try:
        result = call_ai_json(
            "사내 의견을 동일한 정책 대상 단위로 엄격히 재분류한다. 관련성만으로 묶지 않는다. JSON만 출력한다.",
            prompt,
            temperature=0,
        )
        groups=result.get("groups")
        if not isinstance(groups,list) or not groups: raise ValueError()
        refs=[str(ref) for g in groups if isinstance(g,dict) for ref in (g.get("refs") or [])]
        expected={f"c{i}" for i in range(len(rows))}
        if set(refs)!=expected or len(refs)!=len(expected): raise ValueError()
    except Exception:
        return jsonify(error="AI 재분류 결과를 확인하지 못했습니다. 다시 시도해 주세요."),502

    groups=sorted(groups,key=lambda g:len(g.get("refs") or []),reverse=True)
    if len(groups)==1:
        # AI가 한 그룹으로 봐도 이미 존재하는 동일 정책 노드가 있으면 그쪽으로 병합한다.
        candidates=load_policy_candidates(node["parent_id"],{node_id})
        source=[str(groups[0].get("title") or "")]+[rows[int(str(ref)[1:])]["text"] for ref in groups[0].get("refs") or []]
        target=choose_policy_merge_id(source,candidates,{"id":node["parent_id"]} if node["parent_id"] else None)
        if not target: return jsonify(ok=True,groups=1,message="동일한 정책 의견으로 판정되었습니다.")
        merge_targets=[target]
    else:
        # 분리된 각 그룹을 새 노드로 만들기 전에 같은 공지 아래 기존 노드와 다시 비교한다.
        candidates=load_policy_candidates(node["parent_id"],{node_id})
        merge_targets=[]
        reserved=set()
        for g in groups:
            texts=[str(g.get("title") or "")]+[rows[int(str(ref)[1:])]["text"] for ref in g.get("refs") or []]+[str(x) for x in (g.get("keywords") or [])]
            available=[x for x in candidates if str(x["id"]) not in reserved]
            target=choose_policy_merge_id(texts,available,{"id":node["parent_id"]} if node["parent_id"] else None)
            merge_targets.append(target)
            if target: reserved.add(str(target))

    reuse_index=next((i for i,target in enumerate(merge_targets) if not target),None)
    try:
        with db() as c:
            c.execute("BEGIN IMMEDIATE")
            current=[dict(x) for x in c.execute("SELECT co.user_id,u.username author,co.text FROM contributions co JOIN users u ON u.id=co.user_id WHERE co.node_id=? ORDER BY co.created_at,co.user_id",(node_id,)).fetchall()]
            snapshot=lambda xs:[(x["user_id"],x["text"]) for x in xs]
            if snapshot(current)!=snapshot(rows): return jsonify(error="이미 재분류된 의견입니다. 화면을 새로고침했습니다."),409
            total_before=c.execute("SELECT COUNT(*) FROM contributions").fetchone()[0]
            dropped_duplicates=0
            created=0
            merged_existing=0

            for gi,g in enumerate(groups):
                indices=[int(str(ref)[1:]) for ref in g["refs"]]
                first=rows[indices[0]]; title=str(g.get("title") or first["text"])[:40]; category=str(g.get("category") or "기타")[:30]
                keywords=[str(x)[:30] for x in (g.get("keywords") if isinstance(g.get("keywords"),list) else [])[:8]]
                existing_target=merge_targets[gi]
                if existing_target:
                    for idx in indices:
                        dropped_duplicates+=move_contribution(c,node_id,existing_target,rows[idx]["user_id"],rows[idx]["text"])
                    merged_existing+=1
                    continue

                related=[node["parent_id"]] if node["parent_id"] else []
                if gi==reuse_index:
                    c.execute("UPDATE nodes SET text=?,normalized=?,title=?,category=?,keywords=?,related=?,created_by=? WHERE id=?",(first["text"],norm(first["text"]),title,category,json.dumps(keywords,ensure_ascii=False),json.dumps(related),first["user_id"],node_id))
                else:
                    target="n_"+format(int(time.time()*1000)+gi,"x")+secrets.token_hex(2)
                    c.execute("INSERT INTO nodes(id,kind,text,normalized,title,category,keywords,related,parent_id,created_by) VALUES(?,?,?,?,?,?,?,?,?,?)",(target,"opinion",first["text"],norm(first["text"]),title,category,json.dumps(keywords,ensure_ascii=False),json.dumps(related),node["parent_id"],first["user_id"]))
                    created+=1
                    for idx in indices:
                        moved=c.execute("UPDATE contributions SET node_id=? WHERE node_id=? AND user_id=?",(target,node_id,rows[idx]["user_id"]))
                        if moved.rowcount!=1: raise RuntimeError("재분류 중 원본 의견이 변경되었습니다.")

            if reuse_index is None:
                c.execute("DELETE FROM nodes WHERE id=?",(node_id,))
                for r in c.execute("SELECT id,related FROM nodes").fetchall():
                    rel=[x for x in json.loads(r["related"]) if x!=node_id]
                    c.execute("UPDATE nodes SET related=? WHERE id=?",(json.dumps(rel),r["id"]))

            total_after=c.execute("SELECT COUNT(*) FROM contributions").fetchone()[0]
            if total_after!=total_before-dropped_duplicates: raise RuntimeError("재분류 중 의견 수가 변경되었습니다.")
    except (sqlite3.Error,RuntimeError):
        return jsonify(error="재분류 중 데이터가 변경되었습니다. 다시 시도해 주세요."),409
    return jsonify(ok=True,groups=len(groups),merged_existing=merged_existing,created=created,dropped_duplicates=dropped_duplicates)

@app.post("/api/nodes")
@login_required
def create_node():
    d=request.get_json() or {}; text=(d.get("text") or "").strip(); kind=d.get("kind"); a=d.get("analysis") or {}; requested_parent=d.get("parent_id")
    parent_id=requested_parent if kind=="opinion" else None
    if kind not in ("question","opinion") or not text: return jsonify(error="잘못된 입력입니다."),400
    with db() as c:
        if parent_id:
            ptype=c.execute("SELECT kind FROM nodes WHERE id=?",(parent_id,)).fetchone()
            if not ptype or ptype["kind"]!="question": parent_id=None
        if kind=="opinion":
            if parent_id is None:
                exact=c.execute("SELECT id,kind FROM nodes WHERE kind=? AND normalized=? AND parent_id IS NULL",(kind,norm(text))).fetchone()
            else:
                exact=c.execute("SELECT id,kind FROM nodes WHERE kind=? AND normalized=? AND parent_id=?",(kind,norm(text),parent_id)).fetchone()
        else:
            exact=c.execute("SELECT id,kind FROM nodes WHERE kind=? AND normalized=?",(kind,norm(text))).fetchone()
        merge_id=exact["id"] if exact else a.get("merge_id")
        target=c.execute("SELECT id,kind,parent_id FROM nodes WHERE id=?",(merge_id,)).fetchone() if merge_id else None
        if target and kind=="opinion" and target["parent_id"]!=parent_id: target=None
        if target and target["kind"]==kind:
            duplicate=False
            if kind=="opinion":
                duplicate=bool(c.execute("SELECT 1 FROM contributions WHERE node_id=? AND user_id=?",(target["id"],session["user_id"])).fetchone())
                if not duplicate:
                    c.execute("INSERT INTO contributions(node_id,user_id,text) VALUES(?,?,?)",(target["id"],session["user_id"],text))
                    c.execute("UPDATE nodes SET title=?,category=?,keywords=? WHERE id=?",(str(a.get("title") or text)[:40],str(a.get("category") or "기타")[:30],json.dumps((a.get("keywords") if isinstance(a.get("keywords"),list) else [])[:8],ensure_ascii=False),target["id"]))
                    if parent_id:
                        rel=json.loads(c.execute("SELECT related FROM nodes WHERE id=?",(target["id"],)).fetchone()["related"])
                        if parent_id not in rel: c.execute("UPDATE nodes SET related=? WHERE id=?",(json.dumps(([parent_id]+rel)[:4]),target["id"]))
            return jsonify(id=target["id"],merged=True,duplicate=duplicate,parent_linked=bool(parent_id))
        nid="n_"+format(int(time.time()*1000),"x")+secrets.token_hex(2); related=[]
        if parent_id: related.append(parent_id)
        proposed=[] if kind=="question" else [x for x in (a.get("related_ids") if isinstance(a.get("related_ids"),list) else []) if x not in related]
        for rid in proposed:
            other=c.execute("SELECT kind FROM nodes WHERE id=?",(rid,)).fetchone()
            if other and other["kind"]==kind: related.append(rid)
        title=str(a.get("title") or text[:16]); category=str(a.get("category") or "기타"); keywords=a.get("keywords") if isinstance(a.get("keywords"),list) else []
        if parent_id:
            p=c.execute("SELECT title FROM nodes WHERE id=? AND kind='question'",(parent_id,)).fetchone()
            if p and norm(title)==norm(p["title"]): title=text[:24]
        if parent_id and not c.execute("SELECT 1 FROM nodes WHERE id=?",(parent_id,)).fetchone(): parent_id=None
        c.execute("INSERT INTO nodes(id,kind,text,normalized,title,category,keywords,related,parent_id,created_by) VALUES(?,?,?,?,?,?,?,?,?,?)",(nid,kind,text,norm(text),title[:40],category[:30],json.dumps(keywords[:8],ensure_ascii=False),json.dumps(related[:4]),parent_id,session["user_id"]))
        if kind=="opinion": c.execute("INSERT INTO contributions(node_id,user_id,text) VALUES(?,?,?)",(nid,session["user_id"],text))
        return jsonify(id=nid,merged=False,duplicate=False,parent_linked=bool(parent_id))

@app.delete("/api/nodes/<node_id>")
@login_required
def delete_node(node_id):
    admin=bool(session.get("is_admin"))
    with db() as c:
        n=c.execute("SELECT * FROM nodes WHERE id=?",(node_id,)).fetchone()
        if not n: return jsonify(error="노트가 없습니다."),404
        deleted_ids=set()
        if n["kind"]=="opinion":
            if admin:
                c.execute("DELETE FROM nodes WHERE id=?",(node_id,)); deleted_ids.add(node_id)
            else:
                c.execute("DELETE FROM contributions WHERE node_id=? AND user_id=?",(node_id,session["user_id"]))
                if not c.execute("SELECT 1 FROM contributions WHERE node_id=?",(node_id,)).fetchone():
                    c.execute("DELETE FROM nodes WHERE id=?",(node_id,)); deleted_ids.add(node_id)
        elif admin or n["created_by"]==session["user_id"]:
            children=[r["id"] for r in c.execute("SELECT id FROM nodes WHERE kind='opinion' AND parent_id=?",(node_id,)).fetchall()]
            deleted_ids.update(children); deleted_ids.add(node_id)
            c.execute("DELETE FROM nodes WHERE kind='opinion' AND parent_id=?",(node_id,))
            c.execute("DELETE FROM nodes WHERE id=?",(node_id,))
        else: return jsonify(error="작성자만 삭제할 수 있습니다."),403
        if deleted_ids:
            for r in c.execute("SELECT id,related FROM nodes").fetchall():
                rel=[x for x in json.loads(r["related"]) if x not in deleted_ids]
                c.execute("UPDATE nodes SET related=? WHERE id=?",(json.dumps(rel),r["id"]))
    return jsonify(ok=True,deleted_count=len(deleted_ids))

@app.patch("/api/nodes/<node_id>/size")
@login_required
def resize_question(node_id):
    size=(request.get_json() or {}).get("size")
    if size not in ("small","medium","large"): return jsonify(error="크기는 소·중·대만 선택할 수 있습니다."),400
    with db() as c:
        n=c.execute("SELECT kind,created_by FROM nodes WHERE id=?",(node_id,)).fetchone()
        if not n: return jsonify(error="노트가 없습니다."),404
        if n["kind"]!="question" or (not session.get("is_admin") and n["created_by"]!=session["user_id"]): return jsonify(error="공지 작성자 또는 관리자만 크기를 변경할 수 있습니다."),403
        c.execute("UPDATE nodes SET display_size=? WHERE id=?",(size,node_id))
    return jsonify(ok=True,size=size)

@app.delete("/api/connections")
@login_required
def delete_connection():
    d=request.get_json() or {}; a=d.get("a"); b=d.get("b")
    if not a or not b or a==b: return jsonify(error="잘못된 연결입니다."),400
    with db() as c:
        rows=c.execute("SELECT id,kind,parent_id,related FROM nodes WHERE id IN (?,?)",(a,b)).fetchall()
        if len(rows)!=2: return jsonify(error="연결된 노트를 찾을 수 없습니다."),404
        if any(n["kind"]!="opinion" for n in rows): return jsonify(error="공지 노트와의 연결은 해제할 수 없습니다."),403
        if any(n["parent_id"]==(b if n["id"]==a else a) for n in rows): return jsonify(error="공지 노트와의 연결은 해제할 수 없습니다."),403
        for n in rows:
            other=b if n["id"]==a else a
            related=[x for x in json.loads(n["related"]) if x!=other]
            c.execute("UPDATE nodes SET related=? WHERE id=?",(json.dumps(related),n["id"]))
    return jsonify(ok=True)

@app.get("/api/nodes/<node_id>/export.xlsx")
@login_required
def export_large_question(node_id):
    with db() as c:
        q=c.execute("SELECT n.*,u.username author FROM nodes n JOIN users u ON u.id=n.created_by WHERE n.id=?",(node_id,)).fetchone()
        if not q: return jsonify(error="노트가 없습니다."),404
        if q["kind"]!="question" or q["display_size"]!="large": return jsonify(error="대형 공지 노트만 엑셀로 저장할 수 있습니다."),400
        rows=c.execute("""SELECT n.title opinion_title,n.category,u.username author,co.text
          FROM nodes n JOIN contributions co ON co.node_id=n.id JOIN users u ON u.id=co.user_id
          WHERE n.kind='opinion' AND n.parent_id=? ORDER BY n.created_at,co.created_at""",(node_id,)).fetchall()
    wb=Workbook(); ws=wb.active; ws.title="공지 및 의견"
    ws.append(["대형 공지 노트"]); ws.append(["공지 제목",q["title"]]); ws.append(["공지 내용",q["text"]]); ws.append(["공지 작성자",q["author"]]); ws.append(["분류",q["category"]]); ws.append(["키워드",", ".join(json.loads(q["keywords"]))]); ws.append([])
    ws.append(["의견 제목","분류","작성자","개별 의견"])
    for r in rows: ws.append([r["opinion_title"],r["category"],r["author"],r["text"]])
    ws["A1"].font=Font(bold=True,size=16,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor="4F46E5"); ws.merge_cells("A1:D1")
    for cell in ws[8]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="64748B"); cell.alignment=Alignment(horizontal="center")
    for col,width in {"A":24,"B":18,"C":16,"D":60}.items(): ws.column_dimensions[col].width=width
    for row in ws.iter_rows():
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    ws.freeze_panes="A9"; out=BytesIO(); wb.save(out); out.seek(0)
    safe=re.sub(r'[\\/:*?"<>|]+','_',q["title"]).strip() or "공지노트"
    return send_file(out,as_attachment=True,download_name=f"{safe}_의견.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

init_db()
if __name__=="__main__":
    port=int(os.getenv("PORT", "5000"))
    if not IS_RENDER:
        url=f"http://127.0.0.1:{port}"
        threading.Timer(1.0,lambda:webbrowser.open_new(url)).start()
    app.run(host="0.0.0.0",port=port,debug=False,use_reloader=False)
